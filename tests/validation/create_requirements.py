#!/usr/bin/env python3
"""
Script to generate validation_requirements.xlsx for validation tests.
Run this script once to create the Excel file.

Excel format based on varequirementreader.py Mappings:
- Worksheet: "5. Requirements"
- First row: 3 (rows 1-2 are headers)
- Column A: ID
- Column B: Type
- Column C: Description
- Column D: Traces
- Column E: Justification
- Column G: Note
- Column I: Validation Type
"""

import openpyxl
from openpyxl import Workbook


def create_validation_requirements():
    wb = Workbook()

    # Rename default sheet
    ws = wb.active
    ws.title = "5. Requirements"

    # Add headers in rows 1-2 (optional, for readability)
    ws["A1"] = "Requirements"
    ws["A2"] = "ID"
    ws["B2"] = "Type"
    ws["C2"] = "Description"
    ws["D2"] = "Traces"
    ws["E2"] = "Justification"
    ws["G2"] = "Note"
    ws["I2"] = "Validation Type"

    # Add validation requirements starting from row 3
    requirements = [
        {
            "id": "VAL-REQ-001",
            "type": "performance",
            "description": "The system shall process user inputs within 100 milliseconds.",
            "traces": "SYS-001",
            "justification": "User experience",
            "note": "Performance requirement for responsiveness",
            "validation_type": "test",
        },
        {
            "id": "VAL-REQ-002",
            "type": "performance",
            "description": "The system shall support at least 100 concurrent users.",
            "traces": "SYS-002",
            "justification": "System capacity",
            "note": "Scalability requirement",
            "validation_type": "analysis",
        },
        {
            "id": "VAL-REQ-003",
            "type": "functional",
            "description": "The system shall log all user actions for audit purposes.",
            "traces": "SYS-003",
            "justification": "Regulatory compliance",
            "note": "Security and compliance requirement",
            "validation_type": "test",
        },
        {
            "id": "VAL-REQ-004",
            "type": "functional",
            "description": "The system shall not store any user activity data to protect user privacy.",
            "traces": "SYS-004",
            "justification": "Privacy protection",
            "note": "Conflicts with VAL-REQ-003 - logging vs no logging",
            "validation_type": "analysis",
        },
    ]

    for i, req in enumerate(requirements, start=3):
        ws[f"A{i}"] = req["id"]
        ws[f"B{i}"] = req["type"]
        ws[f"C{i}"] = req["description"]
        ws[f"D{i}"] = req["traces"]
        ws[f"E{i}"] = req["justification"]
        ws[f"G{i}"] = req["note"]
        ws[f"I{i}"] = req["validation_type"]

    # Adjust column widths for readability
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["I"].width = 15

    # Save the workbook
    wb.save("validation_requirements.xlsx")
    print("Created validation_requirements.xlsx")


if __name__ == "__main__":
    create_validation_requirements()
